from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


TITLE_COLOR = RGBColor(19, 40, 73)
TEXT_COLOR = RGBColor(49, 60, 74)
ACCENT_BLUE = RGBColor(54, 107, 216)
ACCENT_CYAN = RGBColor(54, 173, 206)
ACCENT_ORANGE = RGBColor(244, 131, 59)
MUTED_TEXT = RGBColor(102, 112, 122)
LIGHT_BG = RGBColor(245, 248, 252)
CARD_BG = RGBColor(255, 255, 255)
LINE_COLOR = RGBColor(219, 228, 237)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_background(slide, title_text, section=None):
    bg = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg.line.fill.background()

    header = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.48)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = TITLE_COLOR
    header.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.45), Inches(0.78), Inches(0.10), Inches(0.62)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT_BLUE
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.70), Inches(0.70), Inches(9.8), Inches(0.7))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    r = p.add_run()
    r.text = title_text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = TITLE_COLOR

    if section:
        sec_box = slide.shapes.add_textbox(Inches(10.4), Inches(0.80), Inches(2.2), Inches(0.4))
        sec_frame = sec_box.text_frame
        sec_p = sec_frame.paragraphs[0]
        sec_p.alignment = PP_ALIGN.RIGHT
        sec_r = sec_p.add_run()
        sec_r.text = section
        sec_r.font.name = "Microsoft YaHei"
        sec_r.font.size = Pt(10)
        sec_r.font.color.rgb = MUTED_TEXT

    footer = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.4), Inches(7.05), Inches(12.53), Inches(0.02)
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = LINE_COLOR
    footer.line.fill.background()


def add_cover(slide):
    bg = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(247, 250, 253)
    bg.line.fill.background()

    top_band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.05)
    )
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = TITLE_COLOR
    top_band.line.fill.background()

    side_block = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.75), Inches(1.45), Inches(3.8), Inches(4.8)
    )
    side_block.fill.solid()
    side_block.fill.fore_color.rgb = LIGHT_BG
    side_block.line.color.rgb = LINE_COLOR

    bar1 = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(8.95), Inches(1.8), Inches(0.15), Inches(3.7)
    )
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = ACCENT_BLUE
    bar1.line.fill.background()

    bar2 = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(9.3), Inches(2.15), Inches(0.15), Inches(3.0)
    )
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = ACCENT_CYAN
    bar2.line.fill.background()

    bar3 = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(9.65), Inches(2.55), Inches(0.15), Inches(2.2)
    )
    bar3.fill.solid()
    bar3.fill.fore_color.rgb = ACCENT_ORANGE
    bar3.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(7.2), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "AI 编程在技术团队中的高效落地"
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = TITLE_COLOR

    sub_box = slide.shapes.add_textbox(Inches(0.82), Inches(3.0), Inches(6.8), Inches(0.8))
    sf = sub_box.text_frame
    sp = sf.paragraphs[0]
    sr = sp.add_run()
    sr.text = "能力边界、使用模式与场景实战"
    sr.font.name = "Microsoft YaHei"
    sr.font.size = Pt(18)
    sr.font.color.rgb = ACCENT_BLUE

    bullets = [
        "面向企业研发、网络攻防、大模型研发",
        "部门内部 AI 编程培训",
        "关键词：效率 / 质量 / 可控",
    ]
    add_bullets(slide, Inches(0.82), Inches(4.0), Inches(6.6), Inches(1.6), bullets, 18)

    meta = slide.shapes.add_textbox(Inches(0.82), Inches(6.3), Inches(5.5), Inches(0.5))
    mf = meta.text_frame
    mp = mf.paragraphs[0]
    mr = mp.add_run()
    mr.text = "讲师 / 部门 / 日期"
    mr.font.name = "Microsoft YaHei"
    mr.font.size = Pt(11)
    mr.font.color.rgb = MUTED_TEXT

    tag_box = slide.shapes.add_textbox(Inches(10.1), Inches(1.8), Inches(2.0), Inches(1.8))
    ttf = tag_box.text_frame
    for idx, text in enumerate(["效率", "质量", "可控"]):
        p = ttf.paragraphs[0] if idx == 0 else ttf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        rr = p.add_run()
        rr.text = text
        rr.font.name = "Microsoft YaHei"
        rr.font.size = Pt(20)
        rr.font.bold = True
        rr.font.color.rgb = TITLE_COLOR if idx < 2 else ACCENT_ORANGE


def add_bullets(slide, left, top, width, height, bullets, font_size=16, color=TEXT_COLOR):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        p.bullet = True
        run = p.add_run()
        run.text = item
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
    return box


def add_quote(slide, text, left=Inches(0.9), top=Inches(5.95), width=Inches(11.5), height=Inches(0.7)):
    quote = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    quote.fill.solid()
    quote.fill.fore_color.rgb = LIGHT_BG
    quote.line.color.rgb = LINE_COLOR
    tf = quote.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = TITLE_COLOR


def add_two_column_cards(slide, left_title, left_bullets, right_title, right_bullets, top=Inches(1.55)):
    card_w = Inches(5.95)
    card_h = Inches(4.7)
    gap = Inches(0.28)
    left = Inches(0.75)
    right = left + card_w + gap
    for x, title, bullets in [
        (left, left_title, left_bullets),
        (right, right_title, right_bullets),
    ]:
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, top, card_w, card_h
        )
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = LINE_COLOR

        head = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, top, card_w, Inches(0.55)
        )
        head.fill.solid()
        head.fill.fore_color.rgb = LIGHT_BG
        head.line.fill.background()

        tbox = slide.shapes.add_textbox(x + Inches(0.22), top + Inches(0.12), card_w - Inches(0.3), Inches(0.3))
        tf = tbox.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = TITLE_COLOR

        add_bullets(slide, x + Inches(0.2), top + Inches(0.75), card_w - Inches(0.4), card_h - Inches(0.95), bullets, 15)


def add_full_card(slide, bullets, top=Inches(1.55), height=Inches(4.95)):
    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), top, Inches(11.85), height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = LINE_COLOR
    add_bullets(slide, Inches(1.0), top + Inches(0.35), Inches(11.1), height - Inches(0.6), bullets, 18)


def add_three_column(slide, columns, top=Inches(1.7), height=Inches(3.9)):
    col_w = Inches(3.8)
    gap = Inches(0.25)
    start = Inches(0.72)
    for idx, (title, bullets, accent) in enumerate(columns):
        x = start + idx * (col_w + gap)
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, top, col_w, height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = LINE_COLOR

        strip = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, top, col_w, Inches(0.12)
        )
        strip.fill.solid()
        strip.fill.fore_color.rgb = accent
        strip.line.fill.background()

        tb = slide.shapes.add_textbox(x + Inches(0.18), top + Inches(0.22), col_w - Inches(0.3), Inches(0.45))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = TITLE_COLOR

        add_bullets(slide, x + Inches(0.14), top + Inches(0.72), col_w - Inches(0.28), height - Inches(0.85), bullets, 14)


def add_four_cards(slide, cards, top=Inches(1.75)):
    positions = [
        (Inches(0.78), top),
        (Inches(6.65), top),
        (Inches(0.78), top + Inches(2.25)),
        (Inches(6.65), top + Inches(2.25)),
    ]
    for (x, y), (title, bullets) in zip(positions, cards):
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(5.85), Inches(2.0)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = LINE_COLOR
        title_box = slide.shapes.add_textbox(x + Inches(0.16), y + Inches(0.15), Inches(5.45), Inches(0.32))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = TITLE_COLOR
        add_bullets(slide, x + Inches(0.14), y + Inches(0.48), Inches(5.5), Inches(1.35), bullets, 13)


def add_formula(slide, text):
    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.8), Inches(11.0), Inches(0.82)
    )
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_BG
    box.line.color.rgb = ACCENT_BLUE
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TITLE_COLOR


def add_timeline(slide, items):
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1.3), Inches(3.4), Inches(10.6), Inches(0.06)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LINE_COLOR
    line.line.fill.background()
    xs = [Inches(2.2), Inches(6.0), Inches(9.8)]
    colors = [ACCENT_CYAN, ACCENT_BLUE, ACCENT_ORANGE]
    for x, (title, subtitle), color in zip(xs, items, colors):
        dot = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, x, Inches(3.18), Inches(0.34), Inches(0.34)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()
        tbox = slide.shapes.add_textbox(x - Inches(0.55), Inches(2.35), Inches(1.5), Inches(0.5))
        tf = tbox.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = TITLE_COLOR
        sbox = slide.shapes.add_textbox(x - Inches(0.95), Inches(3.7), Inches(2.3), Inches(0.45))
        sf = sbox.text_frame
        sp = sf.paragraphs[0]
        sp.alignment = PP_ALIGN.CENTER
        sr = sp.add_run()
        sr.text = subtitle
        sr.font.name = "Microsoft YaHei"
        sr.font.size = Pt(12)
        sr.font.color.rgb = MUTED_TEXT


def add_process_flow(slide, steps):
    start_x = Inches(0.85)
    box_w = Inches(1.85)
    gap = Inches(0.2)
    for idx, step in enumerate(steps):
        x = start_x + idx * (box_w + gap)
        box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, Inches(2.65), box_w, Inches(1.2)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = LINE_COLOR
        num = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, x + Inches(0.68), Inches(2.25), Inches(0.48), Inches(0.48)
        )
        num.fill.solid()
        num.fill.fore_color.rgb = ACCENT_BLUE if idx in (0, 3, 5) else ACCENT_CYAN
        num.line.fill.background()
        nt = slide.shapes.add_textbox(x + Inches(0.81), Inches(2.34), Inches(0.2), Inches(0.2))
        nf = nt.text_frame
        np = nf.paragraphs[0]
        np.alignment = PP_ALIGN.CENTER
        nr = np.add_run()
        nr.text = str(idx + 1)
        nr.font.name = "Microsoft YaHei"
        nr.font.size = Pt(11)
        nr.font.bold = True
        nr.font.color.rgb = RGBColor(255, 255, 255)
        tb = slide.shapes.add_textbox(x + Inches(0.12), Inches(2.95), box_w - Inches(0.24), Inches(0.55))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = step
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = TITLE_COLOR
        if idx < len(steps) - 1:
            arr = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
                x + box_w + Inches(0.02),
                Inches(3.05),
                Inches(0.14),
                Inches(0.28),
            )
            arr.fill.solid()
            arr.fill.fore_color.rgb = LINE_COLOR
            arr.line.fill.background()


def add_prompt_box(slide, title, lines, left, top, width, height):
    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(251, 253, 255)
    card.line.color.rgb = LINE_COLOR
    t = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.12), width - Inches(0.3), Inches(0.3))
    tf = t.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Consolas"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = ACCENT_BLUE
    body = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.45), width - Inches(0.35), height - Inches(0.55))
    bf = body.text_frame
    bf.word_wrap = True
    for idx, line in enumerate(lines):
        p = bf.paragraphs[0] if idx == 0 else bf.add_paragraph()
        p.space_after = Pt(3)
        run = p.add_run()
        run.text = line
        run.font.name = "Consolas"
        run.font.size = Pt(10.5)
        run.font.color.rgb = TEXT_COLOR


def add_checklist(slide, title, bullets, left, top, width):
    title_box = slide.shapes.add_textbox(left, top, width, Inches(0.35))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TITLE_COLOR
    formatted = [f"✓ {bullet}" for bullet in bullets]
    add_bullets(slide, left, top + Inches(0.38), width, Inches(2.2), formatted, 14)


def add_pyramid_rows(slide, rows):
    widths = [Inches(10.6), Inches(9.0), Inches(7.4), Inches(5.8)]
    colors = [ACCENT_BLUE, ACCENT_CYAN, RGBColor(95, 145, 225), ACCENT_ORANGE]
    top = Inches(2.0)
    for idx, (title, bullets) in enumerate(rows):
        width = widths[idx]
        left = (prs.slide_width - width) / 2
        y = top + idx * Inches(0.88)
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, y, width, Inches(0.72)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = colors[idx]
        shape.line.fill.background()
        text = f"{title}：{' / '.join(bullets)}"
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)


def add_route_map(slide, steps, quote):
    add_full_card(slide, [f"{idx + 1}. {text}" for idx, text in enumerate(steps)], top=Inches(1.7), height=Inches(3.75))
    add_quote(slide, quote, top=Inches(5.8), height=Inches(0.8))


slides = []

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_cover(slide)
slides.append(slide)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "为什么现在必须系统谈 AI 编程", "趋势共识")
add_timeline(
    slide,
    [("过去", "代码补全"), ("现在", "协作开发"), ("下一步", "代理执行")],
)
add_quote(slide, "未来效率差距，更多取决于任务拆解、上下文提供、结果校验这三件事。")
slides.append(slide)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "AI 编程到底强在哪，弱在哪", "能力边界")
add_three_column(
    slide,
    [
        ("很强", ["脚本编写", "独立模块开发", "测试代码生成", "文档生成", "重复性重构"], ACCENT_CYAN),
        ("可用但需约束", ["已有规范下的功能迭代", "前后端联调", "存量系统局部改造"], ACCENT_BLUE),
        ("高风险", ["需求本身含糊", "复杂历史系统隐性规则", "缺少验证机制的高风险改动"], ACCENT_ORANGE),
    ],
)
add_quote(slide, "AI 最擅长在边界清楚的任务里高速产出，不擅长替你做业务判断。")
slides.append(slide)

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "为什么大家在企业项目里反而不敢放开用", "真实顾虑")
add_three_column(
    slide,
    [
        ("担心改坏现有功能", ["企业系统牵一发而动全身", "代码背后是流程、权限、数据口径"], ACCENT_ORANGE),
        ("业务理解不完整", ["真实规则不只在文档里", "很多默认共识藏在历史实现里"], ACCENT_BLUE),
        ("代码业务含义不够深", ["很多逻辑是历史妥协", "AI 能解释结构，不一定理解为什么"], ACCENT_CYAN),
    ],
    height=Inches(3.55),
)
add_quote(slide, "企业项目不是不能用 AI，而是不能用“无上下文、无约束、无验证”的 AI。")
slides.append(slide)

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "企业项目里，正确的 AI 编程观", "方法论")
add_full_card(
    slide,
    [
        "从“让 AI 直接写”转向“让 AI 在可控边界内协作”",
        "AI 不是替代业务理解的人，而是放大研发产能的系统",
        "风险越高的任务，越需要更完整上下文、更细粒度拆解、更明确规则、更强验证机制",
    ],
    top=Inches(1.7),
    height=Inches(2.75),
)
add_formula(slide, "AI 编程效果 = 上下文质量 × 任务拆解质量 × 约束质量 × 校验质量")
add_quote(slide, "关键不是让 AI 写得多快，而是让它在你能兜底的范围内写得足够快。")
slides.append(slide)

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "不同角色，要找到不同的 AI 编码模式", "角色打法")
add_four_cards(
    slide,
    [
        ("企业项目研发", ["重点：稳、准、可回归", "最适合：局部代工、测试护栏、最小改动"]),
        ("网络攻防 / 漏洞挖掘", ["重点：脚本效率、分析辅助", "最适合：PoC 样板、解析脚本、日志处理"]),
        ("大模型研发 / 平台工程", ["重点：实验基础设施、评测自动化", "最适合：数据脚本、pipeline glue code、服务封装"]),
        ("非技术同事", ["重点：理解能力边界与协作方式", "最需要知道：AI 快，但不替代责任判断"]),
    ],
)
slides.append(slide)

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "企业研发最值得掌握的 4 种 AI 编码模式", "研发实践")
add_four_cards(
    slide,
    [
        ("新功能脚手架模式", ["页面骨架", "接口定义", "测试样板"]),
        ("局部模块代工模式", ["独立组件", "工具函数", "适配层"]),
        ("带约束的存量修改模式", ["写清楚哪些不能改", "写清楚哪些行为必须保持"]),
        ("测试先行护栏模式", ["先补测试", "再改实现", "适合老系统"]),
    ],
)
add_quote(slide, "企业项目里，AI 写代码前的准备动作，往往比生成动作更重要。", top=Inches(6.15))
slides.append(slide)

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "企业研发的推荐工作流", "安全落地")
add_process_flow(slide, ["明确边界", "补充上下文", "标注禁区", "先出方案", "小步改动", "验证兜底"])
add_quote(slide, "最危险的不是 AI 写得慢，而是把未确认的理解直接变成已提交的代码。")
slides.append(slide)

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "前端实操：从需求到页面骨架", "教程示例")
add_two_column_cards(
    slide,
    "推荐步骤",
    [
        "先拆需求：页面结构、交互行为、状态管理、接口依赖",
        "把项目里的组件规范、路由方式、请求封装喂给 AI",
        "先让 AI 输出页面拆分方案和接口清单",
        "确认后再生成首版代码",
    ],
    "示例提示词",
    [],
)
add_prompt_box(
    slide,
    "Prompt",
    [
        "你现在是这个前端项目的协作开发者。",
        "请基于现有项目规范，完成一个新页面的首版实现方案。",
        "目标页面：风险告警列表页",
        "技术栈：React + TypeScript + Ant Design",
        "要求：优先复用已有组件；先输出页面拆分、状态设计、接口清单；",
        "标出需要我补充的业务规则；我确认后再输出代码。",
    ],
    Inches(6.95),
    Inches(2.15),
    Inches(5.2),
    Inches(3.45),
)
slides.append(slide)

# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "后端实操：新接口 / 老系统改造怎么更安全", "教程示例")
add_two_column_cards(
    slide,
    "新接口开发",
    [
        "先给 AI 明确输入输出",
        "补充实体定义、异常规范、日志要求、幂等要求",
        "先让 AI 出接口设计和边界场景",
        "确认后再生成代码",
    ],
    "老系统改造",
    [
        "先让 AI 解释现有逻辑",
        "识别关键分支和隐性规则",
        "先补测试或回归样例",
        "再做最小改动",
    ],
    top=Inches(1.6),
)
add_prompt_box(
    slide,
    "Prompt",
    [
        "请先不要写代码，先分析这段 service：",
        "1. 用业务语言解释它在做什么",
        "2. 标出关键分支和可能代表的业务规则",
        "3. 识别哪些行为必须保持不变",
        "4. 给出最小改动方案",
        "5. 列出建议补充的测试用例",
    ],
    Inches(0.95),
    Inches(5.75),
    Inches(11.4),
    Inches(0.95),
)
slides.append(slide)

# Slide 11
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "攻防与大模型研发，最适合把 AI 用在哪", "专项场景")
add_two_column_cards(
    slide,
    "网络攻防 / 漏洞挖掘",
    ["PoC / Exp 样板生成", "协议解析脚本", "日志分析与批量数据清洗", "漏洞复盘文档整理"],
    "大模型研发 / 平台工程",
    ["数据处理脚本", "训练 / 推理 pipeline glue code", "评测脚本", "服务封装与部署脚本"],
)
add_quote(slide, "AI 更像高效副驾，擅长加速执行，不替代关键判断。")
slides.append(slide)

# Slide 12
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "提示词、skills、rules，怎么配合使用", "协作机制")
add_three_column(
    slide,
    [
        ("Prompt", ["这次要干什么", "定义目标、上下文、约束、输出格式"], ACCENT_CYAN),
        ("Skills", ["这一类任务通常怎么干", "把高频套路沉淀成可复用方法"], ACCENT_BLUE),
        ("Rules", ["哪些事情永远不能乱干", "修改老逻辑前先解释现状", "高风险改动先补测试"], ACCENT_ORANGE),
    ],
    height=Inches(3.55),
)
add_formula(slide, "好提示词结构 = 背景 + 目标 + 上下文 + 约束 + 输出格式 + 验证要求")
slides.append(slide)

# Slide 13
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "怎么让 AI 更懂你的企业项目", "上下文建设")
add_checklist(
    slide,
    "给 AI 的上下文，至少要包含：",
    [
        "需求背景",
        "业务流程",
        "涉及模块与职责",
        "核心实体 / 字段含义",
        "接口契约",
        "历史限制",
        "不允许破坏的行为",
        "现有测试与验证方式",
    ],
    Inches(1.0),
    Inches(1.8),
    Inches(5.2),
)
context_card = slide.shapes.add_shape(
    MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.9), Inches(5.15), Inches(3.7)
)
context_card.fill.solid()
context_card.fill.fore_color.rgb = LIGHT_BG
context_card.line.color.rgb = LINE_COLOR
tb = slide.shapes.add_textbox(Inches(7.0), Inches(2.25), Inches(4.6), Inches(2.8))
tf = tb.text_frame
for idx, line in enumerate(
    [
        "上下文不是可有可无的补充，",
        "而是企业级 AI 编程的基础设施。",
        "",
        "你给 AI 的企业上下文越完整，",
        "它的输出就越接近“能被采用的工程产出”。",
    ]
):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = line
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(17 if idx in (0, 1) else 14)
    run.font.bold = idx in (0, 1)
    run.font.color.rgb = TITLE_COLOR if idx in (0, 1) else TEXT_COLOR
add_quote(slide, "不要期待 AI 天然理解企业项目，应该主动把企业上下文变成 AI 可消费的材料。")
slides.append(slide)

# Slide 14
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "对团队的真正价值，不只是“写代码更快”", "组织价值")
add_pyramid_rows(
    slide,
    [
        ("效率价值", ["减少样板劳动", "缩短启动时间"]),
        ("质量价值", ["更快补测试", "更快补文档", "覆盖边界场景"]),
        ("知识价值", ["帮新人理解代码", "沉淀可复用方法"]),
        ("组织价值", ["倒逼规范化", "上下文显性化", "流程标准化"]),
    ],
)
add_quote(slide, "AI 编程的长期收益，不只是快，而是把高手经验逐步模板化、制度化。")
slides.append(slide)

# Slide 15
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, "落地建议 + 收尾金句", "总结")
add_route_map(
    slide,
    [
        "从脚本、测试、独立模块开始",
        "每个人先找到自己最顺手的 2-3 个场景",
        "团队沉淀提示词模板、skills、rules",
        "逐步接入需求分析、自测、文档、评审准备",
    ],
    "AI 编程不会替代真正理解业务的人，但会迅速放大那些会定义问题、会约束任务、会验证结果的人。",
)

out_path = "/workspace/AI编程部门内部培训_15页精简版.pptx"
prs.save(out_path)
print(out_path)
