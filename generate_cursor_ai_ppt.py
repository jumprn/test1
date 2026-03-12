from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
BACKGROUND_PATH = ROOT / "cursor-ai-brand-background.png"
OUTPUT_PATH = ROOT / "cursor-ai-best-practices.pptx"
FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BLUE = RGBColor(34, 79, 160)
MAGENTA = RGBColor(217, 45, 127)
RED = RGBColor(214, 71, 71)
ORANGE = RGBColor(242, 149, 54)
GOLD = RGBColor(245, 198, 69)
TITLE = RGBColor(33, 53, 85)
TEXT = RGBColor(53, 68, 86)
MUTED = RGBColor(98, 110, 124)
BORDER = RGBColor(216, 222, 231)
CARD_BG = RGBColor(255, 255, 255)
PROMPT_BG = RGBColor(247, 249, 252)
SOFT_BLUE = RGBColor(235, 242, 252)
SOFT_RED = RGBColor(253, 241, 241)


def rgb_tuple(rgb):
    return (rgb[0], rgb[1], rgb[2], 255)


def load_font(size):
    return ImageFont.truetype(FONT_PATH, size=size)


def blend(c1, c2, t):
    return tuple(int(a + (b - a) * t) for a, b in zip(c1, c2))


def gradient_color(t):
    stops = [
        (0.0, (205, 46, 82, 255)),
        (0.35, (228, 96, 45, 255)),
        (0.62, (246, 205, 58, 255)),
        (1.0, (209, 121, 43, 255)),
    ]
    for idx in range(len(stops) - 1):
        left_t, left = stops[idx]
        right_t, right = stops[idx + 1]
        if left_t <= t <= right_t:
            local = (t - left_t) / (right_t - left_t)
            return blend(left, right, local)
    return stops[-1][1]


def draw_brand_background():
    width, height = 1600, 900
    image = Image.new("RGBA", (width, height), (236, 236, 236, 255))

    warm_overlay = Image.new("RGBA", image.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(warm_overlay, "RGBA")

    draw.ellipse((620, -180, 1030, 170), fill=(235, 116, 62, 86))
    draw.ellipse((820, -120, 1250, 160), fill=(244, 187, 70, 70))
    draw.ellipse((980, -120, 1510, 200), fill=(219, 123, 40, 50))
    draw.polygon([(0, 54), (430, 54), (565, 2), (190, 2)], fill=(240, 104, 63, 50))
    draw.polygon([(1110, 74), (1600, 74), (1600, 5), (1230, 8)], fill=(219, 135, 40, 35))

    warm_overlay = warm_overlay.filter(ImageFilter.GaussianBlur(18))
    image.alpha_composite(warm_overlay)

    stripe = ImageDraw.Draw(image, "RGBA")
    for x in range(40, width - 40):
        t = (x - 40) / (width - 80)
        stripe.line((x, 66, x, 75), fill=gradient_color(t), width=1)

    for offset in range(0, 240, 15):
        stripe.line((width - 315 + offset, 66, width - 260 + offset, 75), fill=(255, 255, 255, 65), width=2)

    brand = ImageDraw.Draw(image, "RGBA")
    blue = (34, 79, 160, 255)
    magenta = (217, 45, 127, 255)
    gray = (160, 165, 171, 255)

    brand.arc((1205, 8, 1256, 56), start=225, end=20, fill=blue, width=5)
    brand.arc((1216, 17, 1245, 44), start=225, end=25, fill=blue, width=4)

    brand.text((1264, 8), "中国电信", font=load_font(28), fill=blue)
    brand.text((1266, 41), "CHINA TELECOM", font=load_font(11), fill=blue)
    brand.line((1452, 12, 1452, 55), fill=gray, width=2)
    brand.text((1474, 7), "翼支付", font=load_font(34), fill=magenta)

    image.save(BACKGROUND_PATH)


def add_background(slide):
    slide.shapes.add_picture(str(BACKGROUND_PATH), 0, 0, width=SLIDE_W, height=SLIDE_H)


def set_text_style(run, size, color, bold=False, name="Microsoft YaHei"):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name


def add_chrome(slide, section, page, footer):
    tag = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.55),
        Inches(0.86),
        Inches(1.75),
        Inches(0.34),
    )
    tag.fill.solid()
    tag.fill.fore_color.rgb = SOFT_BLUE
    tag.line.color.rgb = BLUE
    tf = tag.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = section
    set_text_style(r, 11, BLUE, True)

    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.55),
        Inches(1.28),
        Inches(1.1),
        Inches(0.05),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RED
    line.line.fill.background()

    page_box = slide.shapes.add_textbox(Inches(12.2), Inches(0.88), Inches(0.55), Inches(0.24))
    p = page_box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = page
    set_text_style(r, 11, BLUE, True)

    footer_box = slide.shapes.add_textbox(Inches(0.58), Inches(7.03), Inches(11.9), Inches(0.22))
    p = footer_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = footer
    set_text_style(r, 10, MUTED)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.75), Inches(1.45), Inches(10.9), Inches(0.95))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_text_style(r, 25, TITLE, True)

    if subtitle:
        p2 = tf.add_paragraph()
        p2.space_before = Pt(4)
        r2 = p2.add_run()
        r2.text = subtitle
        set_text_style(r2, 13, MAGENTA, True)


def add_card(slide, x, y, w, h, title, body_lines, accent=BLUE, title_size=15, body_size=14):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = BORDER

    accent_bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, Inches(0.05))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()

    title_box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.14), w - Inches(0.36), Inches(0.28))
    p = title_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_text_style(r, title_size, accent, True)

    body_box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.46), w - Inches(0.34), h - Inches(0.6))
    tf = body_box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for idx, line in enumerate(body_lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        r = p.add_run()
        r.text = f"• {line}"
        set_text_style(r, body_size, TEXT)


def add_plain_panel(slide, x, y, w, h, text, align=PP_ALIGN.LEFT, size=17, color=TEXT, bold=False):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = BORDER
    box = slide.shapes.add_textbox(x + Inches(0.22), y + Inches(0.16), w - Inches(0.44), h - Inches(0.32))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    set_text_style(r, size, color, bold)


def add_quote(slide, text):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.75),
        Inches(5.92),
        Inches(11.85),
        Inches(0.76),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = GOLD
    marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.75), Inches(5.92), Inches(0.12), Inches(0.76))
    marker.fill.solid()
    marker.fill.fore_color.rgb = ORANGE
    marker.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.98), Inches(6.05), Inches(11.4), Inches(0.45))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_text_style(r, 16, TITLE, True)


def add_prompt_box(slide, text, y=Inches(4.1), h=Inches(2.0)):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.75),
        y,
        Inches(11.85),
        h,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PROMPT_BG
    shape.line.color.rgb = BORDER
    marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.75), y, Inches(0.12), h)
    marker.fill.solid()
    marker.fill.fore_color.rgb = BLUE
    marker.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.98), y + Inches(0.14), Inches(11.4), h - Inches(0.26))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_text_style(r, 12.5, TEXT, name="Courier New")


def add_step_row(slide, steps):
    x_positions = [0.75, 3.12, 5.49, 7.86, 10.23]
    for idx, ((title, desc), x) in enumerate(zip(steps, x_positions)):
        add_card(slide, Inches(x), Inches(2.2), Inches(1.92), Inches(1.38), title, [desc], accent=[BLUE, RED, ORANGE, GOLD, MAGENTA][idx], title_size=14, body_size=11)
        if idx < len(steps) - 1:
            arrow_box = slide.shapes.add_textbox(Inches(x + 1.95), Inches(2.72), Inches(0.22), Inches(0.2))
            p = arrow_box.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = "→"
            set_text_style(r, 18, RED, True)


def add_table(slide, headers, rows):
    x_positions = [0.75, 2.55, 6.45]
    widths = [1.55, 3.65, 6.15]
    y = 4.0
    for head, x, w in zip(headers, x_positions, widths):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.42))
        shape.fill.solid()
        shape.fill.fore_color.rgb = SOFT_BLUE
        shape.line.color.rgb = BORDER
        p = shape.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = head
        set_text_style(r, 11, BLUE, True)

    for row_idx, row in enumerate(rows):
        row_y = y + 0.53 + row_idx * 0.54
        for cell, x, w in zip(row, x_positions, widths):
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(row_y), Inches(w), Inches(0.44))
            shape.fill.solid()
            shape.fill.fore_color.rgb = CARD_BG
            shape.line.color.rgb = BORDER
            box = slide.shapes.add_textbox(Inches(x + 0.12), Inches(row_y + 0.08), Inches(w - 0.24), Inches(0.24))
            p = box.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = cell
            set_text_style(r, 10.5, TEXT)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)

    title_box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.72),
        Inches(1.55),
        Inches(7.0),
        Inches(2.65),
    )
    title_box.fill.solid()
    title_box.fill.fore_color.rgb = CARD_BG
    title_box.line.color.rgb = BORDER

    marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.72), Inches(1.55), Inches(0.12), Inches(2.65))
    marker.fill.solid()
    marker.fill.fore_color.rgb = RED
    marker.line.fill.background()

    tf = slide.shapes.add_textbox(Inches(1.0), Inches(1.78), Inches(6.3), Inches(2.1)).text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "Cursor 场景下的 AI 编程最佳实践"
    set_text_style(r1, 17, BLUE, True)

    p2 = tf.add_paragraph()
    p2.space_before = Pt(10)
    r2 = p2.add_run()
    r2.text = "把 AI 从“补全工具”升级为\n可协作的研发搭档"
    set_text_style(r2, 28, TITLE, True)

    p3 = tf.add_paragraph()
    p3.space_before = Pt(10)
    r3 = p3.add_run()
    r3.text = "目标：帮助团队形成一套能稳定产出、可控风险、可复制推广的 AI 编程工作方式。"
    set_text_style(r3, 15, MUTED)

    add_plain_panel(
        slide,
        Inches(8.1),
        Inches(1.7),
        Inches(4.0),
        Inches(2.15),
        "好的 AI 编程结果，通常不是来自“更会问一句话”，而是来自更好的上下文、拆解、约束与验证。",
        align=PP_ALIGN.CENTER,
        size=19,
        color=TITLE,
        bold=True,
    )

    chips = ["需求澄清", "代码探索", "实现与重构", "测试与 Review", "团队规范"]
    chip_x = 0.82
    widths = [1.05, 1.05, 1.25, 1.35, 1.05]
    colors = [BLUE, RED, ORANGE, MAGENTA, GOLD]
    for text, width, color in zip(chips, widths, colors):
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(chip_x),
            Inches(4.55),
            Inches(width),
            Inches(0.42),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = color
        p = shape.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        set_text_style(r, 11, color, True)
        chip_x += width + 0.17

    add_plain_panel(
        slide,
        Inches(0.72),
        Inches(5.2),
        Inches(5.3),
        Inches(1.12),
        "适用对象：研发团队、测试、技术管理者、架构师",
        size=15,
        color=TEXT,
        bold=True,
    )
    add_plain_panel(
        slide,
        Inches(6.28),
        Inches(5.2),
        Inches(5.82),
        Inches(1.12),
        "核心关键词：上下文、拆解、约束、验证、复盘",
        size=15,
        color=TEXT,
        bold=True,
    )

    add_chrome(slide, "内部分享 / AI 编程实践", "01", "建议时长：30-40 分钟")


def build_standard_slide(prs, section, page, footer, title, cards=None, quote=None, prompt=None, steps=None, table=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_chrome(slide, section, page, footer)
    add_title(slide, title)

    if cards:
        for card in cards:
            add_card(slide, *card)
    if steps:
        add_step_row(slide, steps)
    if table:
        add_table(slide, *table)
    if prompt:
        add_prompt_box(slide, prompt["text"], prompt.get("y", Inches(4.05)), prompt.get("h", Inches(2.0)))
    if quote:
        add_quote(slide, quote)


def build_presentation():
    draw_brand_background()

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    title_slide(prs)

    build_standard_slide(
        prs,
        "为什么值得系统化使用",
        "02",
        "关键结论：AI 不是替代工程能力，而是放大工程能力",
        "AI 编程的价值，不只是“写得更快”",
        cards=[
            (Inches(0.75), Inches(2.15), Inches(2.75), Inches(2.35), "1. 降低启动成本", ["新任务、新语言、新模块更容易起步。"], BLUE, 16, 14),
            (Inches(3.62), Inches(2.15), Inches(2.75), Inches(2.35), "2. 扩大能力边界", ["一个人可以兼顾实现、测试、文档、脚本。"], RED, 16, 14),
            (Inches(6.49), Inches(2.15), Inches(2.75), Inches(2.35), "3. 提升交付吞吐", ["尤其适合样板代码、重构、批量修改、测试补齐。"], ORANGE, 16, 14),
            (Inches(9.36), Inches(2.15), Inches(2.75), Inches(2.35), "4. 倒逼工程规范", ["上下文清晰、测试充分、结构良好的仓库更容易被 AI 利用。"], GOLD, 16, 14),
        ],
        quote="核心判断：AI 最适合“信息量大、模式重复、需要快速试错”的研发活动。",
    )

    build_standard_slide(
        prs,
        "Cursor 的定位",
        "03",
        "角色切换越清晰，产出越稳定",
        "在 Cursor 里，建议把 AI 当成 4 类角色使用",
        cards=[
            (Inches(0.75), Inches(2.2), Inches(5.55), Inches(1.55), "1. 探索员", ["快速扫描代码库、定位入口、解释调用链、总结模块职责。"], BLUE, 16, 14),
            (Inches(6.55), Inches(2.2), Inches(5.55), Inches(1.55), "2. 实现助手", ["按明确边界完成函数、组件、接口、迁移脚本、测试代码。"], RED, 16, 14),
            (Inches(0.75), Inches(4.0), Inches(5.55), Inches(1.55), "3. 评审员", ["对 diff 做风险扫描，找潜在 bug、回归点、边界条件、缺失测试。"], ORANGE, 16, 14),
            (Inches(6.55), Inches(4.0), Inches(5.55), Inches(1.55), "4. 教练", ["解释陌生代码、给出重构思路、比较方案、补齐知识盲区。"], MAGENTA, 16, 14),
        ],
        quote="更高效的方式：让 AI 参与需求理解、方案拆解、实现、验证和复盘的完整链路。",
    )

    build_standard_slide(
        prs,
        "最佳实践总原则",
        "04",
        "经验：AI 编程的上限，常常取决于人的任务设计能力",
        "先给原则，再谈技巧",
        cards=[
            (Inches(0.75), Inches(2.1), Inches(3.72), Inches(1.45), "原则 1：先对齐任务", ["明确目标、范围、约束、验收标准。"], BLUE, 14, 12.5),
            (Inches(4.59), Inches(2.1), Inches(3.72), Inches(1.45), "原则 2：优先给上下文", ["让 AI 看到正确文件、接口、报错和已有模式。"], RED, 14, 12.5),
            (Inches(8.43), Inches(2.1), Inches(3.72), Inches(1.45), "原则 3：任务拆小", ["一次只解决一个问题，改完立刻验证。"], ORANGE, 14, 12.5),
            (Inches(0.75), Inches(3.85), Inches(3.72), Inches(1.45), "原则 4：让 AI 解释修改", ["要求列出改动点、影响面、未覆盖风险。"], MAGENTA, 14, 12.5),
            (Inches(4.59), Inches(3.85), Inches(3.72), Inches(1.45), "原则 5：人负责验收", ["AI 可以生成，但不拥有上线决策权。"], GOLD, 14, 12.5),
            (Inches(8.43), Inches(3.85), Inches(3.72), Inches(1.45), "原则 6：沉淀团队资产", ["把高质量 Prompt 模板化、规范化、流程化。"], BLUE, 14, 12.5),
        ],
        quote="先分析、再实现、再验证，是最稳的 AI 协作节奏。",
    )

    build_standard_slide(
        prs,
        "工作流 1",
        "05",
        "先分析再实现，能显著减少“写出来但不对”的返工",
        "从模糊需求到可执行任务",
        steps=[
            ("目标", "我要解决什么业务问题？"),
            ("范围", "改哪些模块？不改哪些模块？"),
            ("约束", "语言、接口、兼容性、性能要求。"),
            ("验收", "测试、页面、日志、输出标准。"),
            ("行动", "先分析，再提方案，再实现。"),
        ],
        prompt={
            "text": "建议先这样问 Cursor：\n\n"
                    "“你先不要写代码。请基于当前仓库帮我完成 3 件事：\n"
                    "1. 找出和这个需求最相关的文件与调用链；\n"
                    "2. 给出最小改动方案；\n"
                    "3. 列出可能影响的测试点和风险点。\n\n"
                    "需求目标：...\n改动范围：...\n不要改动：...\n验收标准：...”",
            "y": Inches(4.15),
            "h": Inches(1.9),
        },
    )

    build_standard_slide(
        prs,
        "工作流 2",
        "06",
        "正确上下文 > 华丽提示词",
        "优先让 AI “读代码”，再让 AI “写代码”",
        cards=[
            (Inches(0.75), Inches(2.15), Inches(5.65), Inches(3.2), "推荐做法", [
                "先让 Cursor 总结模块职责、入口函数、关键数据流。",
                "要求引用具体文件和函数，而不是泛泛解释。",
                "让 AI 用现有代码风格完成实现，而不是凭空创造新模式。",
                "要求它说明为什么选这个修改点而不是别处。",
            ], BLUE, 16, 13),
            (Inches(6.55), Inches(2.15), Inches(5.65), Inches(3.2), "常见问题", [
                "没有给上下文，AI 按通用经验瞎猜。",
                "一次贴太多无关文件，核心信息被稀释。",
                "没要求它引用证据，回答听起来合理但脱离仓库实际。",
                "直接让它重构大片代码，最后人无法 Review。",
            ], RED, 16, 13),
        ],
        quote="一个非常有效的习惯：先问“仓库里现在是怎么做的”，再问“这次我应该怎么改”。",
    )

    build_standard_slide(
        prs,
        "工作流 3",
        "07",
        "AI 越强，越要控制改动粒度",
        "用“小步快跑”替代“一次生成一大坨”",
        cards=[
            (Inches(0.75), Inches(2.1), Inches(3.7), Inches(1.45), "步骤 A", ["先让 AI 只做方案和影响分析，确认改动边界。"], BLUE, 15, 12.5),
            (Inches(4.8), Inches(2.1), Inches(3.7), Inches(1.45), "步骤 B", ["一次只改一个文件或一个子问题，方便人类 Review。"], RED, 15, 12.5),
            (Inches(8.85), Inches(2.1), Inches(3.35), Inches(1.45), "步骤 C", ["每轮修改后立刻跑测试、lint、编译或关键路径验证。"], ORANGE, 15, 12.5),
        ],
        table=(
            ["场景", "不推荐", "更推荐"],
            [
                ("功能开发", "“把整个需求都实现了”", "“先补接口层，再补 service，再补测试”"),
                ("Bug 修复", "“帮我修掉这个问题”", "“先定位根因，再给最小修复 diff”"),
                ("重构", "一次性批量重写", "“分批迁移，每批可回滚、有测试覆盖”"),
            ],
        ),
    )

    build_standard_slide(
        prs,
        "工作流 4",
        "08",
        "如果你不知道怎么验证，通常也不该让 AI 直接开始写",
        "让 AI 先写“验证方式”，再写实现",
        steps=[
            ("验证", "先设计测试或验证清单。"),
            ("覆盖", "正常路径、边界、异常、回归。"),
            ("定位", "确认现有代码为什么失败。"),
            ("方案", "给出最小修复方案。"),
            ("落地", "修复后补齐测试。"),
        ],
        prompt={
            "text": "可复用提示词：\n\n"
                    "“请先不要修改生产代码，先基于当前实现补充或设计测试用例：\n"
                    "- 正常路径\n- 边界条件\n- 异常分支\n- 回归风险\n\n"
                    "然后告诉我：\n1. 现有代码为什么会失败；\n2. 最小修复方案是什么；\n3. 修复后需要补哪些测试。”",
            "y": Inches(4.12),
            "h": Inches(1.95),
        },
    )

    build_standard_slide(
        prs,
        "Prompt 模板",
        "09",
        "提示词不是玄学，模板化后就是团队方法论",
        "3 个最常用、最值得团队复用的 Prompt 模板",
        cards=[
            (Inches(0.75), Inches(2.15), Inches(3.7), Inches(2.15), "模板 1：代码探索", ["请阅读相关文件，告诉我入口、调用链、关键数据结构、改动建议，并引用具体文件位置。"], BLUE, 15, 13),
            (Inches(4.78), Inches(2.15), Inches(3.7), Inches(2.15), "模板 2：最小改动实现", ["基于现有风格完成最小改动，不额外引入新框架，修改后总结影响面并列出验证步骤。"], RED, 15, 13),
            (Inches(8.81), Inches(2.15), Inches(3.39), Inches(2.15), "模板 3：Review / 风险扫描", ["请站在 reviewer 视角检查这次 diff，优先找 bug、回归风险、缺失测试。"], ORANGE, 15, 13),
        ],
        quote="模板设计建议：固定 5 个槽位——目标 / 上下文 / 约束 / 输出格式 / 验收标准。",
    )

    build_standard_slide(
        prs,
        "上下文管理",
        "10",
        "一个经验：越复杂的需求，越要像写需求文档一样写给 AI",
        "在 Cursor 中，真正拉开差距的是上下文管理",
        cards=[
            (Inches(0.75), Inches(2.1), Inches(5.65), Inches(3.2), "高质量上下文应该包含", [
                "当前目标和背景：为什么要改。",
                "相关文件与错误信息：而不是只贴结论。",
                "已有实现模式：类似模块、历史方案、命名规范。",
                "边界和禁区：哪些地方不能改、不能破坏什么。",
                "验证标准：测试、运行、页面、日志、性能。",
            ], BLUE, 16, 13),
            (Inches(6.55), Inches(2.1), Inches(5.65), Inches(3.2), "实操建议", [
                "长对话中及时重置上下文，避免历史噪音污染当前任务。",
                "复杂任务把分析、实现、Review 分成多个对话阶段。",
                "发现 AI 偏题时，先补上下文或缩小任务。",
                "对关键结论要求引用代码证据和文件位置。",
            ], RED, 16, 13),
        ],
    )

    build_standard_slide(
        prs,
        "Review 与风险控制",
        "11",
        "优秀团队不是“少写错”，而是“更快发现错”",
        "把 AI 用在 Review，通常是投入产出比最高的环节",
        cards=[
            (Inches(0.75), Inches(2.1), Inches(3.75), Inches(1.5), "检查 bug", ["空指针、未处理异常、并发问题、状态遗漏、边界条件漏掉。"], RED, 15, 12.5),
            (Inches(4.85), Inches(2.1), Inches(3.75), Inches(1.5), "检查回归", ["公共接口是否变更、老逻辑是否被破坏、兼容性是否受影响。"], ORANGE, 15, 12.5),
            (Inches(8.95), Inches(2.1), Inches(3.25), Inches(1.5), "检查测试", ["关键路径是否覆盖、失败场景是否覆盖、断言是否足够严格。"], BLUE, 15, 12.5),
        ],
        prompt={
            "text": "Review 提示词建议：\n\n"
                    "“请用 code review 的方式审查这次改动：\n"
                    "- 优先指出可能的 bug、风险和行为变化；\n"
                    "- 标注影响文件或逻辑点；\n"
                    "- 判断是否缺失测试；\n"
                    "- 如果没有明显问题，也请说明仍然存在的测试盲区。\n\n"
                    "不要只复述改动内容。”",
            "y": Inches(4.05),
            "h": Inches(1.95),
        },
    )

    build_standard_slide(
        prs,
        "团队治理",
        "12",
        "规模化效果 = 工具能力 × 工程规范 × 团队习惯",
        "如果要规模化落地，需要建立 4 个团队机制",
        cards=[
            (Inches(0.75), Inches(2.2), Inches(5.55), Inches(1.55), "机制 1：提示词模板库", ["沉淀探索、实现、测试、Review、文档等常用模板。"], BLUE, 15, 13),
            (Inches(6.55), Inches(2.2), Inches(5.55), Inches(1.55), "机制 2：AI 使用边界", ["明确哪些代码、数据、密钥、客户信息不能放进上下文。"], RED, 15, 13),
            (Inches(0.75), Inches(4.05), Inches(5.55), Inches(1.55), "机制 3：验收标准", ["AI 生成代码必须过测试、lint、编译、Review，不得跳过工程门禁。"], ORANGE, 15, 13),
            (Inches(6.55), Inches(4.05), Inches(5.55), Inches(1.55), "机制 4：案例复盘", ["定期复盘高价值用例和失败案例，让方法论持续进化。"], MAGENTA, 15, 13),
        ],
        quote="团队采用 AI 编程，不是“装个工具”就结束了，而是一次研发工作流升级。",
    )

    build_standard_slide(
        prs,
        "常见误区",
        "13",
        "结论：AI 编程失败，通常不是因为模型不够强，而是流程设计不到位",
        "最容易让团队“用了 AI 但没得到收益”的 6 个误区",
        cards=[
            (Inches(0.75), Inches(2.05), Inches(3.72), Inches(1.42), "误区 1", ["把 AI 当搜索引擎，问泛问题，不给仓库上下文。"], RED, 14, 12.5),
            (Inches(4.59), Inches(2.05), Inches(3.72), Inches(1.42), "误区 2", ["一次性让 AI 改太多，导致 Review 和回滚成本过高。"], ORANGE, 14, 12.5),
            (Inches(8.43), Inches(2.05), Inches(3.72), Inches(1.42), "误区 3", ["只看它能不能生成，不看是否可验证、可维护。"], GOLD, 14, 12.5),
            (Inches(0.75), Inches(3.8), Inches(3.72), Inches(1.42), "误区 4", ["把 AI 输出直接当事实，不要求引用仓库证据。"], MAGENTA, 14, 12.5),
            (Inches(4.59), Inches(3.8), Inches(3.72), Inches(1.42), "误区 5", ["没有沉淀复用模板，每个人都从零摸索。"], BLUE, 14, 12.5),
            (Inches(8.43), Inches(3.8), Inches(3.72), Inches(1.42), "误区 6", ["忽略安全边界，把不该暴露的信息交给 AI。"], RED, 14, 12.5),
        ],
        quote="AI 编程的价值，最终要靠可验证的结果来证明。",
    )

    build_standard_slide(
        prs,
        "建议现场 Demo",
        "14",
        "最佳 Demo：一个真实小任务胜过十个空泛技巧",
        "分享时，推荐演示 1 个完整闭环",
        steps=[
            ("Step 1", "给出一个真实需求或 bug。"),
            ("Step 2", "先让 Cursor 分析代码与定位影响点。"),
            ("Step 3", "再让它提出最小改动方案。"),
            ("Step 4", "分步实现，并跑验证命令。"),
            ("Step 5", "最后让 AI 站在 reviewer 视角复查。"),
        ],
        quote="不要演示“AI 一次写了很多代码”，而要演示“AI 如何帮助你更快做出正确决策、减少试错、补齐验证”。",
    )

    build_standard_slide(
        prs,
        "落地建议",
        "15",
        "谢谢 / 可按公司案例继续替换内容",
        "结论与落地建议",
        cards=[
            (Inches(0.75), Inches(2.1), Inches(5.65), Inches(2.9), "今天希望带走的 3 个观点", [
                "AI 编程的关键不是“写得快”，而是“更快得到正确结果”。",
                "Cursor 最好的用法，是把 AI 放进完整研发链路，而不是只做补全。",
                "团队收益来自模板、流程、验证和复盘，而不是个人技巧。",
            ], BLUE, 16, 13),
            (Inches(6.55), Inches(2.1), Inches(5.65), Inches(2.9), "建议下一步行动", [
                "挑 2-3 个真实场景做试点：Bug 修复、测试补齐、批量重构。",
                "沉淀一套团队 Prompt 模板与安全边界。",
                "每周复盘“最好用”和“最翻车”的案例。",
            ], RED, 16, 13),
        ],
        quote="Q&A：我们不是在学习一个新工具，而是在学习一种新的研发协作方式。",
    )

    prs.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_presentation()
    print(f"Generated {OUTPUT_PATH.name}")
    print(f"Generated {BACKGROUND_PATH.name}")
