from __future__ import annotations

from pathlib import Path
from textwrap import dedent

import matplotlib
from matplotlib import font_manager

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parent
DOCX_PATH = ROOT / "一种面向Web服务的自适应闭环安全扫描方法及系统_专利交底书.docx"
FIG1_PATH = ROOT / "附图1_自适应闭环扫描总体流程.png"
FIG2_PATH = ROOT / "附图2_资产语义建模与修复回放流程.png"


def configure_plot_font() -> str:
    candidates = [
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            font_manager.fontManager.addfont(str(path))
            font_name = font_manager.FontProperties(fname=str(path)).get_name()
            plt.rcParams["font.family"] = font_name
            plt.rcParams["axes.unicode_minus"] = False
            return font_name
    plt.rcParams["axes.unicode_minus"] = False
    return "DejaVu Sans"


def set_chinese_font(run, font_name: str = "SimSun", size: int | None = None, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold


def add_paragraph(document: Document, text: str, *, style: str | None = None, bold: bool = False, size: int = 12):
    paragraph = document.add_paragraph(style=style)
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(6)
    paragraph_format.line_spacing = 1.35
    run = paragraph.add_run(text)
    set_chinese_font(run, size=size, bold=bold)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Normal"]
    run = paragraph.add_run(text)
    set_chinese_font(run, size=16 if level == 1 else 14, bold=True)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def build_figure(path: Path, title: str, boxes: list[tuple[float, float, str, str]], arrows: list[tuple[int, int]]) -> None:
    fig, ax = plt.subplots(figsize=(14, 8))
    ax.set_xlim(0, 14)
    ax.set_ylim(0, 8)
    ax.axis("off")

    palette = {
        "entry": "#D6EAF8",
        "core": "#D5F5E3",
        "risk": "#FADBD8",
        "verify": "#FCF3CF",
        "control": "#E8DAEF",
    }

    centers: list[tuple[float, float]] = []
    for x, y, text, kind in boxes:
        width, height = 2.7, 1.1
        patch = FancyBboxPatch(
            (x - width / 2, y - height / 2),
            width,
            height,
            boxstyle="round,pad=0.05,rounding_size=0.08",
            linewidth=1.5,
            edgecolor="#34495E",
            facecolor=palette[kind],
        )
        ax.add_patch(patch)
        ax.text(x, y, text, ha="center", va="center", fontsize=11, wrap=True)
        centers.append((x, y))

    for start_idx, end_idx in arrows:
        x1, y1 = centers[start_idx]
        x2, y2 = centers[end_idx]
        ax.annotate(
            "",
            xy=(x2, y2),
            xytext=(x1, y1),
            arrowprops=dict(arrowstyle="->", lw=1.6, color="#2C3E50", shrinkA=28, shrinkB=28),
        )

    ax.text(7, 7.5, title, ha="center", va="center", fontsize=16, weight="bold")
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)


def generate_figures() -> None:
    configure_plot_font()
    build_figure(
        FIG1_PATH,
        "附图1 自适应闭环安全扫描总体流程",
        [
            (1.7, 5.7, "待测Web服务\n接入层", "entry"),
            (4.3, 5.7, "资产语义指纹\n构建模块", "core"),
            (7.0, 5.7, "令牌化交互沙箱\n生成模块", "core"),
            (9.7, 5.7, "分层自适应扫描\n执行模块", "risk"),
            (12.2, 5.7, "漏洞证据\n融合模块", "verify"),
            (12.2, 2.4, "修复回放验证\n模块", "verify"),
            (9.2, 2.4, "策略编排与\n优先级更新", "control"),
            (5.7, 2.4, "知识图谱与\n攻击路径推理", "control"),
            (2.5, 2.4, "结果输出与\n处置建议", "entry"),
        ],
        [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (5, 6), (6, 7), (7, 3), (6, 8)],
    )

    build_figure(
        FIG2_PATH,
        "附图2 资产语义建模与修复回放流程",
        [
            (2.0, 5.6, "HTTP/HTTPS响应\n采集", "entry"),
            (4.6, 5.6, "前端脚本、接口、\n鉴权元素解析", "core"),
            (7.1, 5.6, "语义特征向量与\n服务画像生成", "core"),
            (9.8, 5.6, "疑似漏洞触发\n链路构造", "risk"),
            (12.1, 5.6, "多源证据\n置信聚合", "verify"),
            (10.0, 2.4, "修复补丁上线后\n自动回放", "verify"),
            (6.6, 2.4, "差异比对与\n误报剔除", "control"),
            (3.0, 2.4, "扫描策略库\n增量更新", "control"),
        ],
        [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (5, 6), (6, 7), (7, 3)],
    )


def patent_sections() -> list[tuple[str, str]]:
    return [
        (
            "一、初拟的发明名称",
            "一种面向Web服务的自适应闭环安全扫描方法及系统",
        ),
        (
            "二、名词解释",
            dedent(
                """
                Web服务：通过HTTP或HTTPS协议向浏览器、移动端或系统间调用方提供页面、接口、文件或控制能力的软件服务实体。
                资产语义指纹（Semantic Asset Fingerprint）：对目标Web服务的页面结构、接口模式、脚本行为、认证方式、组件特征以及业务语义进行联合编码后形成的高维特征描述。
                令牌化交互沙箱（Tokenized Interaction Sandbox）：一种受控执行环境，用于在不破坏生产服务稳定性的前提下，向目标服务注入经过约束的测试令牌、请求序列和模拟用户行为，以触发潜在漏洞路径。
                漏洞触发链（Vulnerability Trigger Chain）：从入口页面、参数源、鉴权状态、调用顺序到目标敏感动作的完整触发路径集合。
                多源证据融合（Multi-source Evidence Fusion）：对响应报文、DOM变化、日志片段、回显差异、状态码漂移和时序行为进行联合判定，以提高漏洞识别可信度的过程。
                修复回放验证（Remediation Replay Verification）：在目标系统完成补丁修复后，自动重放历史触发链并比对结果，确认漏洞是否真正消除以及是否产生回归问题。
                风险知识图谱（Risk Knowledge Graph）：用于描述资产、组件、接口、参数、权限、数据流和漏洞类型之间关联关系的图谱化模型。
                扫描策略编排器（Scan Strategy Orchestrator）：根据资产语义指纹和历史扫描结果动态选择、排序并调整扫描插件、探针强度和验证路径的控制模块。
                """
            ).strip(),
        ),
        (
            "三、所属技术领域",
            "本发明涉及网络安全检测、Web应用安全、自动化漏洞挖掘以及安全运维技术领域，特别是涉及一种适用于复杂Web服务的自适应闭环安全扫描方法及系统。",
        ),
        (
            "四、背景技术",
            dedent(
                """
                1、现有技术方案描述：当前针对Web服务的安全扫描方案主要包括以下几类：
                （1）基于规则库的被动扫描方案：通过对URL、端口、Banner、页面标题、常见路径或已知漏洞特征进行匹配，识别服务类型并执行预定义探测。该方案实现成本较低，但对于前后端分离、动态渲染、接口链复杂和权限状态强相关的现代Web服务，往往只能发现表层问题。
                （2）传统爬虫结合漏洞插件的主动扫描方案：利用爬虫遍历站点页面，抽取表单和接口，再调用SQL注入、跨站脚本、越权访问等插件进行测试。该方案在静态站点中有效，但在单页应用、复杂登录态、多步骤业务流和异步接口场景下，容易出现链路中断、漏扫和误报。
                （3）人工渗透测试辅助方案：由安全工程师根据经验构造测试用例、切换角色、观察页面逻辑和接口行为，综合判断系统风险。该方法深度较高，但依赖人工经验，成本高、覆盖不稳定且难以持续化执行。

                2、现有技术存在的问题：
                （1）上下文缺失：现有扫描器通常把页面、接口和参数视为离散对象，缺乏对业务流程、认证状态和跨页面调用关系的整体理解，难以发现依赖时序和状态切换的漏洞。
                （2）策略静态：多数扫描规则在任务启动前一次性确定，无法根据目标服务的实时反馈动态调整探测深度，导致对复杂系统扫描不足、对简单系统又产生冗余探测。
                （3）误报与漏报并存：仅基于单次回显、状态码或关键字命中做判断，容易把正常业务提示误判为漏洞，同时又遗漏需要多源证据交叉验证的真实风险。
                （4）缺少闭环：现有方案大多止步于“发现问题”，很少在漏洞修复后自动重放原始触发链，无法验证修复是否彻底，也难以将经验反哺到后续扫描策略中。
                （5）适配性不足：对于微服务、CDN前置、GraphQL、WebSocket、前端加密参数、验证码旁路和多租户后台等新型Web服务形态，传统扫描框架适配成本高。

                3、本发明最主要解决的技术问题：
                如何构建一种面向复杂Web服务的安全扫描机制，使其能够先理解目标服务的语义结构和业务上下文，再动态生成适配性的扫描路径、通过多源证据确认漏洞可信度，并在修复后自动回放验证，从而形成“识别—扫描—验证—学习”的逻辑闭环。
                """
            ).strip(),
        ),
        (
            "五、发明创造的目的",
            dedent(
                """
                主要流程步骤：
                步骤一（主要发明点）：采集目标Web服务的页面、接口、脚本、会话和组件信息，构建能够表达业务上下文、状态切换和暴露面的资产语义指纹。
                步骤二：依据资产语义指纹，自动生成令牌化交互沙箱与漏洞触发链，对需要登录、切换角色、跨页面跳转或多步提交的业务流进行可控重放。
                步骤三：执行分层自适应扫描，对不同路径采用不同强度和不同插件组合，并在扫描过程中根据返回结果实时调整策略。
                步骤四：对扫描所得结果进行多源证据融合，输出高置信漏洞结论、攻击路径说明和修复建议。
                步骤五（另一主要发明点）：在系统修复后，对已发现风险的触发链执行自动回放验证，并将验证结果反哺到风险知识图谱与扫描策略编排器中。

                主要发明点在于步骤一、步骤二、步骤五，即通过“语义理解 + 受控交互沙箱 + 修复回放学习”的方式，使扫描器不再只做静态探测，而是具备对业务流的理解能力、对扫描路径的动态编排能力以及对修复效果的持续确认能力。

                总结性描述：
                本发明旨在解决传统Web扫描器对现代复杂Web服务理解不足、误报漏报较多以及缺少修复验证闭环的问题。通过本发明，可以在不显著增加人工参与的前提下，提高对复杂业务漏洞的发现深度、降低误报率，并建立漏洞发现与修复验证相互反馈的持续安全运营机制。
                """
            ).strip(),
        ),
        (
            "六、发明创造的技术方案以及具体实施例",
            dedent(
                """
                1、使用场景：
                本方案涉及一个包含以下实体的系统：待测Web服务、扫描控制台、资产语义建模模块、令牌化交互沙箱、扫描策略编排器、多源证据融合模块、风险知识图谱、修复回放验证模块以及结果输出终端。扫描控制台可部署在企业内部安全平台、托管式扫描平台或私有化漏洞管理平台中。待测Web服务可以是门户网站、管理后台、微服务网关、单页应用、移动端后端接口、SaaS管理平台或其他基于HTTP/HTTPS开放能力的业务系统。

                2、具体实施例：
                优选实施例（结合附图1的流程图）：
                （1）系统初始化：
                管理员在扫描控制台中配置目标资产域名、入口URL、允许访问的角色测试账号、频率阈值、黑白名单路径以及修复验证开关。系统加载基础插件库、业务流模板、组件漏洞库和风险知识图谱。

                （2）资产语义指纹构建：
                触发条件：用户发起一次针对目标Web服务的扫描任务。
                执行动作：系统通过浏览器仿真内核和HTTP探针同时访问目标站点，采集首页、登录页、静态资源、JS脚本、API描述、接口响应头、Cookie策略、重定向链和前端路由信息。
                处理动作：资产语义建模模块对采集内容进行解析，提取服务框架指纹、前端组件库、参数命名规律、鉴权因子、页面跳转关系、接口依赖和表单语义标签，并生成资产语义指纹。
                得到结果：形成关于“该Web服务如何暴露能力、如何鉴权、如何传参、如何流转业务状态”的结构化画像。

                （3）令牌化交互沙箱生成：
                触发条件：系统发现目标服务存在登录、验证码旁路、角色切换、工单提交、审批流或订单流等多步骤业务链。
                执行动作：令牌化交互沙箱模块根据资产语义指纹创建受控会话环境，注入测试账户、时效令牌、伪造但可追踪的业务标识和可回收测试数据。
                处理动作：系统把页面动作、接口调用和状态转移组织为漏洞触发链。例如在越权检测中，系统构造“普通用户登录—访问列表接口—提取对象ID—替换高权限对象ID—提交查看请求—观察回显差异”的链路；在存储型跨站脚本检测中，系统构造“低风险标签注入—编码变换—展示页回显—脚本执行监控”的链路。
                得到结果：得到可控、可追踪、可重放的业务交互路径，为后续扫描提供上下文基础。

                （4）分层自适应扫描：
                触发条件：完成目标画像和触发链构造后，扫描策略编排器开始选择插件与探针强度。
                执行动作：编排器根据资产语义指纹和历史知识图谱，给不同接口、参数和页面分配不同的扫描策略；对高风险管理入口优先执行鉴权绕过、越权、逻辑缺陷和敏感信息泄露探测，对普通展示页优先执行XSS、文件遍历和配置泄露探测。
                处理动作：在扫描执行过程中，系统实时分析响应时延、状态码变化、页面DOM变动、脚本执行轨迹、服务端错误片段和回显模式。如果发现目标服务进入新的业务状态、暴露隐藏接口或返回具有高风险特征的异常内容，系统会立即提升该路径的扫描权重，动态追加二次验证探针。
                得到结果：实现“先粗扫定位、后深扫确认”的自适应扫描过程，兼顾效率与深度。

                （5）多源证据融合与漏洞确认：
                触发条件：某条触发链出现异常响应、敏感信息外泄、权限边界变化或页面执行异常。
                执行动作：多源证据融合模块同时分析请求/响应差异、DOM结构变化、浏览器控制台事件、埋点日志、接口返回对象规模、鉴权标识变化和同一触发链多次执行的一致性。
                处理动作：系统为每类证据分配可信权重，并利用风险知识图谱检查该目标组件是否存在已知弱点、该接口是否处于高价值业务节点、该参数是否控制敏感对象范围。
                得到结果：仅当综合置信度达到设定阈值时，系统生成正式漏洞结论，并输出漏洞名称、影响范围、复现步骤、原始请求、响应证据和修复建议。

                （6）修复回放验证：
                触发条件：目标系统完成修复并由运维或开发在平台中标记“待复测”。
                执行动作：修复回放验证模块调取历史漏洞对应的触发链、测试令牌和证据基线，在相同或等价条件下自动重放。
                处理动作：系统将新结果与历史风险证据进行差异比对，判断漏洞是否被彻底消除、是否仅隐藏了报错信息但本质风险仍然存在，或者是否因补丁产生新的可用性/权限回归。
                得到结果：输出“修复成功”“修复不彻底”或“产生回归风险”的验证结论，并将验证结果回写到风险知识图谱和扫描策略编排器中。

                拓展技术方案：
                a. 在优选实施例基础上，可支持跨域微服务联合扫描。系统通过API网关日志、前端路由和统一鉴权信息，将多个后端服务映射到同一业务图谱中，实现跨服务业务流漏洞检测。
                b. 可支持基于大模型辅助的语义页面理解，用于识别复杂页面中的“审批”“导出”“重置口令”“切换租户”等关键动作，从而提高业务逻辑漏洞检测能力。
                c. 可支持面向不同租户或不同环境（测试、预发、生产影子）的隔离扫描策略，使扫描强度与令牌权限动态匹配，降低对真实业务的影响。
                d. 可支持把修复回放结果反哺到企业SDL流程，在代码合并或上线前自动触发风险回归检查。

                3、有益效果：
                （1）理解上下文：通过资产语义指纹与风险知识图谱，本发明能够理解Web服务的业务结构与状态切换关系，显著提高对复杂业务漏洞的识别能力。
                （2）动态适配：通过扫描策略编排器与令牌化交互沙箱，系统可根据目标反馈实时调整探测路径与探针强度，提升扫描效率并减少无效请求。
                （3）降低误报：通过多源证据融合和一致性校验，不再依赖单一回显或状态码判断，从而降低误报率并提升结果可解释性。
                （4）形成闭环：通过修复回放验证模块，本发明将漏洞发现、修复确认和知识沉淀连接起来，形成可持续学习的安全扫描闭环。
                （5）部署灵活：既可作为独立扫描平台部署，也可与CI/CD、漏洞管理平台、工单系统和安全运营中心对接，适用性强。
                """
            ).strip(),
        ),
        (
            "七、发明人认为要保护的发明内容的技术要点以及相应的有益效果",
            dedent(
                """
                技术要点一：一种针对Web服务构建资产语义指纹的方法，其将页面结构、接口行为、组件特征、鉴权方式和业务流关系进行联合编码，形成可用于扫描策略选择的目标画像。
                有益效果：使扫描系统能够理解目标Web服务的上下文和结构特征，为复杂场景下的准确扫描提供基础。

                技术要点二：一种基于资产语义指纹自动生成令牌化交互沙箱与漏洞触发链的方法。
                有益效果：使扫描器具备在受控环境中模拟真实业务流的能力，可发现传统静态规则难以触达的多步骤、强状态依赖型漏洞。

                技术要点三：一种在扫描执行过程中根据实时响应特征动态调整探针强度和插件组合的自适应扫描机制。
                有益效果：能够在控制扫描成本的同时，把资源聚焦到高风险路径上，提高漏洞挖掘效率。

                技术要点四：一种结合响应差异、DOM变更、控制台事件、日志片段和历史知识图谱进行漏洞可信度判定的多源证据融合方法。
                有益效果：提升漏洞结论的准确性和可解释性，降低误报和漏报。

                技术要点五：一种对历史漏洞触发链进行自动回放并将结果反哺到知识图谱与策略编排器中的修复验证闭环机制。
                有益效果：确保漏洞修复具有可验证性，并让扫描系统随着使用过程不断优化，具备持续演进能力。
                """
            ).strip(),
        ),
        (
            "八、附图",
            "附图1：自适应闭环安全扫描总体流程图。\n附图2：资产语义建模与修复回放流程图。",
        ),
        (
            "九、其他",
            dedent(
                """
                对于理解交底书中的技术方案有帮助的专利/论文/期刊，可进一步参考如下方向资料：
                1. Web应用漏洞扫描与自动化验证相关论文；
                2. 基于知识图谱的攻击路径分析研究；
                3. Browser automation、headless crawling、API security testing等相关期刊与技术报告；
                4. 业务逻辑漏洞检测、越权检测、修复回归验证等相关专利或公开资料。

                补充说明：
                本发明并不限于特定的漏洞类型，可扩展应用于SQL注入、跨站脚本、越权访问、敏感数据泄露、逻辑缺陷、SSRF、文件上传、接口滥用、工作流绕过等Web服务风险场景；亦不限于特定厂商组件或部署形态。
                """
            ).strip(),
        ),
    ]


def create_document() -> None:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "SimSun"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    styles["Normal"].font.size = Pt(12)

    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("专利交底书\n一种面向Web服务的自适应闭环安全扫描方法及系统")
    set_chinese_font(run, font_name="SimHei", size=18, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subrun = subtitle.add_run("（根据用户给定模板扩展生成）")
    set_chinese_font(subrun, size=11)

    add_paragraph(
        document,
        "说明：本交底书突出“创意性、新颖性、技术闭环”三项特征，重点创新点为资产语义理解、令牌化交互沙箱、自适应扫描编排及修复回放验证。",
        bold=False,
        size=11,
    )

    for heading, body in patent_sections():
        add_heading(document, heading)
        for block in body.split("\n"):
            add_paragraph(document, block.strip())

    add_heading(document, "附图说明与流程摘要")
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    headers = ["附图编号", "流程模块", "说明"]
    for idx, text in enumerate(headers):
        paragraph = hdr_cells[idx].paragraphs[0]
        run = paragraph.add_run(text)
        set_chinese_font(run, bold=True)
        shade_cell(hdr_cells[idx], "D9EAF7")

    figure_rows = [
        ("附图1", "总体闭环", "展示从目标接入、语义建模、自适应扫描、证据融合到修复回放学习的全生命周期闭环。"),
        ("附图2", "局部细化", "展示资产语义建模、触发链构造、证据确认、回放复测和策略增量更新之间的细粒度协同关系。"),
    ]
    for code, module, desc in figure_rows:
        row = table.add_row().cells
        for cell, text in zip(row, (code, module, desc)):
            p = cell.paragraphs[0]
            r = p.add_run(text)
            set_chinese_font(r)

    document.add_paragraph()
    for img_path, caption in (
        (FIG1_PATH, "附图1 自适应闭环安全扫描总体流程图"),
        (FIG2_PATH, "附图2 资产语义建模与修复回放流程图"),
    ):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img_path), width=Inches(6.5))
        cp = document.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption)
        set_chinese_font(cr, size=11, bold=True)

    document.add_section(WD_SECTION.NEW_PAGE)
    add_heading(document, "可选权利要求撰写思路（供后续专利代理人整理）")
    claims = [
        "一种面向Web服务的自适应闭环安全扫描方法，其特征在于，包括：构建资产语义指纹、生成令牌化交互沙箱、形成漏洞触发链、执行分层自适应扫描、进行多源证据融合以及实施修复回放验证。",
        "根据权利要求1所述的方法，其中资产语义指纹至少包含页面结构特征、接口调用关系、组件特征、鉴权方式特征及业务流状态转移特征。",
        "根据权利要求1所述的方法，其中扫描策略编排器基于实时响应结果动态更新探针强度、插件组合和验证优先级。",
        "根据权利要求1所述的方法，其中修复回放验证结果用于反向更新风险知识图谱和扫描策略库。",
        "一种实现上述方法的系统，包括资产语义建模模块、令牌化交互沙箱模块、扫描策略编排器、多源证据融合模块、修复回放验证模块和结果输出模块。",
    ]
    for idx, claim in enumerate(claims, start=1):
        add_paragraph(document, f"{idx}. {claim}")

    document.save(DOCX_PATH)


def main() -> None:
    generate_figures()
    create_document()
    print(f"Generated: {DOCX_PATH}")
    print(f"Generated: {FIG1_PATH}")
    print(f"Generated: {FIG2_PATH}")


if __name__ == "__main__":
    main()
