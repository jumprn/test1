from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm as PptCm
from pptx.util import Pt as PptPt


ROOT = Path(__file__).resolve().parent
FIG1_PATH = ROOT / "附图1_自适应闭环扫描总体流程.png"
FIG2_PATH = ROOT / "附图2_资产语义建模与修复回放流程.png"
DOCX_PATH = ROOT / "Web服务扫描专利_PPT汇报文案.docx"
PPTX_PATH = ROOT / "Web服务扫描专利汇报稿.pptx"


SLIDES = [
    {
        "page_no": "封面/目录页",
        "title": "汇报提纲",
        "subtitle": "围绕“背景、摘要、系统方法、新颖性、创造性、实用性”六部分展开",
        "bullets": [
            "1. 专利背景",
            "2. 专利摘要",
            "3. 系统方法说明",
            "4. 新颖性说明",
            "5. 创造性说明",
            "6. 实用性说明",
        ],
        "speaker_notes": [
            "本次汇报基于《一种面向Web服务的自适应闭环安全扫描方法及系统》展开。",
            "重点突出方案的技术闭环、创新点以及在企业安全运营中的落地价值。",
        ],
    },
    {
        "page_no": "专利介绍1",
        "title": "背景",
        "blocks": [
            (
                "项目来源",
                [
                    "来源于企业级Web服务安全治理、攻防演练和持续漏洞运营场景。",
                    "面向门户网站、管理后台、微服务接口、单页应用和SaaS平台的自动化扫描需求。",
                ],
            ),
            (
                "技术领域",
                [
                    "网络安全检测、Web应用安全、自动化漏洞挖掘与安全运维。",
                    "尤其适用于复杂业务流、强状态依赖和多角色鉴权的Web服务环境。",
                ],
            ),
            (
                "应用场景",
                [
                    "日常漏洞巡检、重大活动前专项排查、上线前安全验收、修复后回归验证。",
                    "可与漏洞管理平台、CI/CD流水线、安全运营中心联动使用。",
                ],
            ),
            (
                "解决的问题",
                [
                    "传统扫描器对现代Web服务业务链理解不足，容易漏扫越权、逻辑缺陷等深层漏洞。",
                    "仅依赖单次回显或状态码判断，误报较多，且缺少修复后自动复测机制。",
                ],
            ),
        ],
    },
    {
        "page_no": "专利介绍2",
        "title": "摘要",
        "blocks": [
            (
                "核心思想",
                [
                    "先理解目标Web服务，再决定如何扫描、如何确认以及如何验证修复。",
                    "通过“语义建模 + 自适应扫描 + 回放验证”形成完整技术闭环。",
                ],
            ),
            (
                "方案摘要",
                [
                    "系统首先采集页面、接口、脚本、鉴权和业务流信息，构建资产语义指纹。",
                    "随后自动生成令牌化交互沙箱和漏洞触发链，对复杂业务路径实施分层自适应扫描。",
                    "最后通过多源证据融合输出高置信结论，并在修复后自动回放验证，实现持续学习。",
                ],
            ),
            (
                "预期效果",
                [
                    "提高复杂业务漏洞发现能力，降低误报漏报。",
                    "将漏洞发现、修复确认和知识沉淀统一到同一平台流程中。",
                ],
            ),
        ],
    },
    {
        "page_no": "专利介绍3-1",
        "title": "系统方法说明",
        "blocks": [
            (
                "系统主要包括",
                [
                    "资产语义指纹构建模块：采集页面、接口、脚本、鉴权及组件特征并生成服务画像。",
                    "令牌化交互沙箱模块：构建受控会话、测试令牌和可追踪业务数据。",
                    "扫描策略编排模块：根据目标画像和历史结果动态选择扫描路径与探针强度。",
                    "分层自适应扫描执行模块：对高风险页面、接口和对象执行差异化扫描。",
                    "多源证据融合模块：综合请求响应、DOM变化、日志和一致性结果判定漏洞。",
                    "风险知识图谱模块：维护资产、组件、漏洞类型及攻击路径之间的关联。",
                    "修复回放验证模块：在补丁上线后重放历史触发链，确认漏洞是否真正消除。",
                ],
            ),
            (
                "系统输出",
                [
                    "输出漏洞名称、影响范围、复现请求、证据链、修复建议和复测结论。",
                    "支持向漏洞管理平台、工单系统和持续交付平台同步结果。",
                ],
            ),
        ],
    },
    {
        "page_no": "专利介绍3-2",
        "title": "系统方法说明（流程）",
        "blocks": [
            (
                "系统分析流程",
                [
                    "步骤1：采集目标Web服务的页面、接口、路由、鉴权和组件信息，生成资产语义指纹。",
                    "步骤2：构造令牌化交互沙箱与漏洞触发链，执行分层自适应扫描并完成多源证据确认。",
                    "步骤3：在修复后自动回放历史触发链，输出修复成功/不彻底/存在回归风险的结论。",
                ],
            ),
            (
                "流程关键点A",
                [
                    "理解业务上下文，不再只对URL和参数做静态探测。",
                    "对越权、逻辑缺陷、状态切换类问题可实施更精准检测。",
                ],
            ),
            (
                "流程关键点B",
                [
                    "扫描结果可以反向更新知识图谱和策略库，实现越用越准。",
                    "形成“识别—扫描—验证—学习”的闭环安全运营体系。",
                ],
            ),
        ],
        "figure": str(FIG1_PATH),
    },
    {
        "page_no": "专利介绍4",
        "title": "新颖性说明",
        "blocks": [
            (
                "检索与对比思路",
                [
                    "可与传统规则扫描、爬虫式扫描、人工渗透辅助方案进行对比。",
                    "重点关注现有技术是否具备对业务语义理解、动态策略编排和修复回放闭环的统一实现。",
                ],
            ),
            (
                "本方案的新颖点",
                [
                    "提出资产语义指纹，将页面结构、接口关系、鉴权状态和业务流特征联合建模。",
                    "提出令牌化交互沙箱，用受控方式模拟真实用户操作与多角色切换流程。",
                    "提出修复回放验证机制，使扫描系统具备漏洞修复后的自动复测和知识反哺能力。",
                ],
            ),
            (
                "相对现有技术区别",
                [
                    "现有技术通常止步于漏洞发现，本方案进一步覆盖漏洞确认和修复验证。",
                    "现有技术多基于静态规则，本方案可根据目标反馈实时调整扫描强度和路径。",
                ],
            ),
        ],
    },
    {
        "page_no": "专利介绍5",
        "title": "创造性说明",
        "blocks": [
            (
                "相对现有技术的优势",
                [
                    "把目标理解、路径构造、漏洞确认和修复验证连接为一体，不是简单的插件堆叠。",
                    "能够覆盖多步骤、强状态依赖、跨角色切换的复杂业务漏洞场景。",
                ],
            ),
            (
                "主要创造性体现",
                [
                    "创造性一：将“资产语义指纹”作为扫描策略输入，使扫描行为具备上下文感知能力。",
                    "创造性二：将“令牌化交互沙箱”与“漏洞触发链”结合，使复杂业务动作可控、可追踪、可回放。",
                    "创造性三：将“多源证据融合”与“修复回放验证”并入同一技术链路，形成持续学习闭环。",
                ],
            ),
            (
                "技术效果",
                [
                    "提高深层漏洞发现能力，降低误报，提升结果可解释性。",
                    "使系统不仅能发现问题，还能验证修复是否有效，具备显著的实质性特点和进步。",
                ],
            ),
        ],
    },
    {
        "page_no": "专利介绍6",
        "title": "实用性说明",
        "blocks": [
            (
                "可应用场景",
                [
                    "企业门户、运维后台、客户服务平台、政企SaaS系统、API网关、移动端后端服务。",
                    "适合日常巡检、等保整改、重大保障、上线前安全验收、修复后回归测试等场景。",
                ],
            ),
            (
                "落地方式",
                [
                    "可独立部署为Web安全扫描平台，也可与现有漏洞管理、工单和CI/CD系统集成。",
                    "支持按租户、按项目、按环境进行策略隔离和分级扫描。",
                ],
            ),
            (
                "可能价值",
                [
                    "提升安全团队对复杂Web资产的自动化治理能力，降低人工渗透成本。",
                    "缩短漏洞发现到修复确认的周期，为企业持续安全运营提供可量化支撑。",
                ],
            ),
        ],
        "figure": str(FIG2_PATH),
    },
]


def set_run_font(run, name: str = "SimSun", size: int = 12, bold: bool = False, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        from docx.shared import RGBColor

        run.font.color.rgb = RGBColor.from_string(color)


def build_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    document.styles["Normal"].font.name = "SimSun"
    document.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    document.styles["Normal"].font.size = Pt(12)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Web服务扫描专利PPT汇报文案")
    set_run_font(run, name="SimHei", size=18, bold=True, color="1F4E79")

    intro = document.add_paragraph()
    intro_run = intro.add_run(
        "说明：本文件按照“汇报提纲—背景—摘要—系统方法说明—新颖性说明—创造性说明—实用性说明”的PPT逻辑组织，可直接复制到PPT页面。"
    )
    set_run_font(intro_run, size=11)

    for idx, slide in enumerate(SLIDES, start=1):
        heading = document.add_paragraph()
        hr = heading.add_run(f"第{idx}页：{slide['title']}")
        set_run_font(hr, name="SimHei", size=15, bold=True, color="C55A11")

        if slide.get("subtitle"):
            sub = document.add_paragraph()
            sr = sub.add_run(f"页面定位：{slide['subtitle']}")
            set_run_font(sr, size=11, color="555555")

        if slide.get("bullets"):
            for bullet in slide["bullets"]:
                p = document.add_paragraph(style=None)
                p.paragraph_format.left_indent = Cm(0.7)
                r = p.add_run(f"• {bullet}")
                set_run_font(r, size=12, bold=True, color="1F3A93")

        for title_text, points in slide.get("blocks", []):
            p = document.add_paragraph()
            r = p.add_run(f"【{title_text}】")
            set_run_font(r, name="SimHei", size=13, bold=True, color="1F3A93")
            for point in points:
                item = document.add_paragraph()
                item.paragraph_format.left_indent = Cm(0.8)
                ir = item.add_run(f"• {point}")
                set_run_font(ir, size=12)

        if slide.get("speaker_notes"):
            note_title = document.add_paragraph()
            nr = note_title.add_run("【讲解提示】")
            set_run_font(nr, name="SimHei", size=12, bold=True, color="00838F")
            for note in slide["speaker_notes"]:
                np = document.add_paragraph()
                np.paragraph_format.left_indent = Cm(0.8)
                npr = np.add_run(f"• {note}")
                set_run_font(npr, size=11, color="444444")

        document.add_paragraph()

    document.save(DOCX_PATH)


def add_top_banner(slide, title: str, subtitle: str | None = None) -> None:
    blue = RGBColor(29, 78, 216)
    orange = RGBColor(234, 88, 12)
    dark_blue = RGBColor(22, 56, 122)
    light_bg = RGBColor(250, 250, 252)

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = light_bg

    top_bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, PptCm(33.87), PptCm(1.2))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = blue
    top_bar.line.color.rgb = blue

    logo_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, PptCm(0.3), PptCm(0.15), PptCm(4.8), PptCm(1.9))
    logo_box.fill.solid()
    logo_box.fill.fore_color.rgb = dark_blue
    logo_box.line.color.rgb = dark_blue
    tf = logo_box.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = "专利汇报"
    r1.font.name = "SimHei"
    r1.font.size = PptPt(18)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(255, 255, 255)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Web服务安全扫描"
    r2.font.name = "SimSun"
    r2.font.size = PptPt(9)
    r2.font.color.rgb = RGBColor(220, 230, 255)

    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, PptCm(0), PptCm(1.9), PptCm(33.87), PptCm(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = orange
    line.line.color.rgb = orange

    title_box = slide.shapes.add_textbox(PptCm(9.0), PptCm(0.25), PptCm(18), PptCm(1.6))
    title_tf = title_box.text_frame
    title_tf.clear()
    tp = title_tf.paragraphs[0]
    tp.alignment = PP_ALIGN.CENTER
    tr = tp.add_run()
    tr.text = title
    tr.font.name = "SimHei"
    tr.font.size = PptPt(26)
    tr.font.color.rgb = orange
    tr.font.bold = False

    if subtitle:
        sub_box = slide.shapes.add_textbox(PptCm(23.5), PptCm(1.2), PptCm(9), PptCm(0.6))
        sub_tf = sub_box.text_frame
        sub_tf.clear()
        sp = sub_tf.paragraphs[0]
        sp.alignment = PP_ALIGN.RIGHT
        sr = sp.add_run()
        sr.text = subtitle
        sr.font.name = "SimSun"
        sr.font.size = PptPt(9)
        sr.font.color.rgb = RGBColor(110, 110, 110)


def add_text_block(slide, left: float, top: float, width: float, height: float, title: str, points: list[str]) -> None:
    box = slide.shapes.add_textbox(PptCm(left), PptCm(top), PptCm(width), PptCm(height))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = f"❖ {title}"
    r.font.name = "SimHei"
    r.font.size = PptPt(18)
    r.font.color.rgb = RGBColor(31, 58, 147)
    r.font.bold = True
    for point in points:
        item = tf.add_paragraph()
        item.level = 1
        item.space_before = 0
        item.space_after = 0
        ir = item.add_run()
        ir.text = point
        ir.font.name = "SimSun"
        ir.font.size = PptPt(14)
        ir.font.color.rgb = RGBColor(22, 56, 122)


def add_content_frame(slide) -> None:
    border = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, PptCm(1.3), PptCm(2.5), PptCm(28.8), PptCm(15.0))
    border.fill.background()
    border.line.color.rgb = RGBColor(70, 70, 70)
    border.line.width = PptPt(1)


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = PptCm(33.867)
    prs.slide_height = PptCm(19.05)
    blank_layout = prs.slide_layouts[6]

    # 1. Outline
    slide = prs.slides.add_slide(blank_layout)
    add_top_banner(slide, "汇报提纲")
    y = 3.6
    for idx, text in enumerate(SLIDES[0]["bullets"], start=1):
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.HEXAGON, PptCm(2.7), PptCm(y - 0.4), PptCm(2.2), PptCm(1.7))
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(14, 165, 233) if idx % 2 else RGBColor(147, 197, 253)
        circle.line.color.rgb = RGBColor(107, 114, 128)
        tf = circle.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(idx)
        r.font.name = "SimHei"
        r.font.size = PptPt(20)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

        line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, PptCm(5.2), PptCm(y + 0.45), PptCm(18), PptCm(0.05))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(40, 40, 40)
        line.line.color.rgb = RGBColor(40, 40, 40)

        text_box = slide.shapes.add_textbox(PptCm(7.6), PptCm(y - 0.05), PptCm(12), PptCm(0.8))
        ttf = text_box.text_frame
        ttf.clear()
        tp = ttf.paragraphs[0]
        tr = tp.add_run()
        tr.text = text.split(". ", 1)[1]
        tr.font.name = "SimHei"
        tr.font.size = PptPt(24)
        tr.font.bold = True
        tr.font.color.rgb = RGBColor(31, 58, 147)
        y += 2.05

    # 2..8 detail pages
    for slide_data in SLIDES[1:]:
        slide = prs.slides.add_slide(blank_layout)
        add_top_banner(slide, f"{slide_data['page_no']}：{slide_data['title']}")
        add_content_frame(slide)

        if slide_data["title"] == "系统方法说明（流程）":
            title_box = slide.shapes.add_textbox(PptCm(1.7), PptCm(2.8), PptCm(6.5), PptCm(0.8))
            tf = title_box.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "系统分析流程："
            r.font.name = "SimHei"
            r.font.size = PptPt(18)
            r.font.color.rgb = RGBColor(31, 58, 147)
            r.font.bold = True

            if FIG1_PATH.exists():
                slide.shapes.add_picture(str(FIG1_PATH), PptCm(1.9), PptCm(3.6), width=PptCm(14.2))

            add_text_block(slide, 17.5, 4.0, 10.7, 4.0, "步骤1-3", slide_data["blocks"][0][1])
            add_text_block(slide, 17.5, 9.5, 10.7, 3.2, "关键点A", slide_data["blocks"][1][1])
            add_text_block(slide, 17.5, 13.0, 10.7, 3.2, "关键点B", slide_data["blocks"][2][1])
            continue

        if slide_data["title"] == "实用性说明" and FIG2_PATH.exists():
            add_text_block(slide, 1.7, 3.0, 13.2, 6.1, slide_data["blocks"][0][0], slide_data["blocks"][0][1])
            add_text_block(slide, 1.7, 9.1, 13.2, 3.6, slide_data["blocks"][1][0], slide_data["blocks"][1][1])
            add_text_block(slide, 1.7, 12.7, 13.2, 3.6, slide_data["blocks"][2][0], slide_data["blocks"][2][1])
            slide.shapes.add_picture(str(FIG2_PATH), PptCm(16.0), PptCm(4.2), width=PptCm(12.8))
            continue

        y = 3.0
        for title_text, points in slide_data.get("blocks", []):
            height = 2.2 + 0.72 * len(points)
            add_text_block(slide, 1.7, y, 27.2, height, title_text, points)
            y += height + 0.45

    prs.save(PPTX_PATH)


def main() -> None:
    build_docx()
    build_pptx()
    print(f"Generated: {DOCX_PATH}")
    print(f"Generated: {PPTX_PATH}")


if __name__ == "__main__":
    main()
