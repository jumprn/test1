#!/usr/bin/env python3

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]

FILES = [
    (
        ROOT / "docs" / "ai-coding-advocacy-formal.md",
        ROOT / "docs" / "ai-coding-advocacy-formal.docx",
    ),
    (
        ROOT / "docs" / "ai-coding-advocacy-brief.md",
        ROOT / "docs" / "ai-coding-advocacy-brief.docx",
    ),
]


def configure_document(document: Document) -> None:
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(11)


def add_markdown_line(document: Document, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        document.add_paragraph("")
        return

    heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
    if heading_match:
        level = min(len(heading_match.group(1)), 4)
        document.add_heading(heading_match.group(2).strip(), level=level)
        return

    bullet_match = re.match(r"^-\s+(.*)$", stripped)
    if bullet_match:
        document.add_paragraph(bullet_match.group(1).strip(), style="List Bullet")
        return

    number_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
    if number_match:
        document.add_paragraph(number_match.group(2).strip(), style="List Number")
        return

    quote_match = re.match(r"^>\s?(.*)$", stripped)
    if quote_match:
        paragraph = document.add_paragraph(style="Intense Quote")
        paragraph.add_run(quote_match.group(1).strip())
        return

    document.add_paragraph(stripped)


def convert_markdown_to_docx(source: Path, target: Path) -> None:
    document = Document()
    configure_document(document)

    for line in source.read_text(encoding="utf-8").splitlines():
        add_markdown_line(document, line)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def main() -> None:
    for source, target in FILES:
        convert_markdown_to_docx(source, target)
        print(f"Generated {target.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
